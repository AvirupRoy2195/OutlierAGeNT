"""LLM-assisted outlier agent.

The agent can run fully offline with rule-based planning. If the OpenAI Python
package and OPENAI_API_KEY are available, it can ask an LLM to select the
method and imputation strategy, then validates that choice before execution.
"""

from __future__ import annotations

import argparse
import json
import os
from dataclasses import dataclass
from itertools import zip_longest
from pathlib import Path
from typing import Any

from outlier import OutlierTool, load_csv, save_csv


CONTEXT_PATH = Path(__file__).with_name("context.json")
SKILLS_PATH = Path(__file__).with_name("skills.md")
ENV_PATH = Path(__file__).with_name(".env")


@dataclass(frozen=True)
class Plan:
    method: str
    imputation: str
    domain: str
    reason: str
    source: str

    def to_dict(self) -> dict[str, str]:
        return {
            "method": self.method,
            "imputation": self.imputation,
            "domain": self.domain,
            "reason": self.reason,
            "source": self.source,
        }


class OutlierAgent:
    def __init__(
        self,
        *,
        context_path: Path = CONTEXT_PATH,
        skills_path: Path = SKILLS_PATH,
        env_path: Path = ENV_PATH,
        use_llm: bool = True,
    ) -> None:
        load_env_file(env_path)
        self.context = json.loads(context_path.read_text(encoding="utf-8"))
        self.skills = skills_path.read_text(encoding="utf-8")
        self.use_llm = use_llm
        self.tool = OutlierTool()

    def run(
        self,
        records: list[dict[str, Any]],
        *,
        columns: list[str] | None = None,
        domain: str | None = None,
        method: str | None = None,
        imputation: str | None = None,
        task: str | None = None,
    ) -> dict[str, Any]:
        plan = self.plan(
            records,
            columns=columns,
            domain=domain,
            method=method,
            imputation=imputation,
            task=task,
        )
        defaults = self.context["defaults"]
        cleaned_records, report = self.tool.detect_and_impute(
            records,
            columns=columns,
            method=plan.method,
            imputation=plan.imputation,
            iqr_multiplier=float(defaults["iqr_multiplier"]),
            zscore_threshold=float(defaults["zscore_threshold"]),
            modified_zscore_threshold=float(defaults["modified_zscore_threshold"]),
        )
        return {
            "plan": plan.to_dict(),
            "report": report.to_dict(),
            "cleaned_records": cleaned_records,
        }

    def plan(
        self,
        records: list[dict[str, Any]],
        *,
        columns: list[str] | None = None,
        domain: str | None = None,
        method: str | None = None,
        imputation: str | None = None,
        task: str | None = None,
    ) -> Plan:
        explicit_method = method.lower().strip() if method else None
        explicit_imputation = imputation.lower().strip() if imputation else None

        if explicit_method and explicit_imputation:
            return self._validated_plan(
                {
                    "method": explicit_method,
                    "imputation": explicit_imputation,
                    "domain": domain or "user-specified",
                    "reason": "User supplied both method and imputation.",
                    "source": "user",
                }
            )

        llm_plan = None
        if self.use_llm:
            llm_plan = self._plan_with_llm(
                records,
                columns=columns,
                domain=domain,
                method=explicit_method,
                imputation=explicit_imputation,
                task=task,
            )
        if llm_plan:
            return llm_plan

        return self._plan_with_rules(
            records,
            columns=columns,
            domain=domain,
            method=explicit_method,
            imputation=explicit_imputation,
        )

    def _plan_with_rules(
        self,
        records: list[dict[str, Any]],
        *,
        columns: list[str] | None,
        domain: str | None,
        method: str | None,
        imputation: str | None,
    ) -> Plan:
        detected_domain = (domain or self._infer_domain(records, columns)).lower()
        profiles = self.context["domain_profiles"]
        profile = profiles.get(detected_domain, profiles["general"])
        selected_method = method or profile["method"]
        selected_imputation = imputation or profile["imputation"]

        if len(records) < 20 and not method:
            selected_method = "modified_zscore"

        return self._validated_plan(
            {
                "method": selected_method,
                "imputation": selected_imputation,
                "domain": detected_domain,
                "reason": profile.get("reason", "Selected from local context rules."),
                "source": "rules",
            }
        )

    def _plan_with_llm(
        self,
        records: list[dict[str, Any]],
        *,
        columns: list[str] | None,
        domain: str | None,
        method: str | None,
        imputation: str | None,
        task: str | None,
    ) -> Plan | None:
        api_key = os.getenv("OUTLIER_AGENT_API_KEY") or os.getenv("OPENAI_API_KEY")
        endpoint = (
            os.getenv("OUTLIER_AGENT_ENDPOINT")
            or os.getenv("OPENAI_BASE_URL")
            or self.context["llm"].get("endpoint")
        )
        model = os.getenv("OUTLIER_AGENT_MODEL") or self.context["llm"]["model"]
        temperature = float(
            os.getenv("OUTLIER_AGENT_TEMPERATURE", self.context["llm"]["temperature"])
        )
        if not configured_secret(api_key):
            return None

        try:
            from openai import OpenAI
        except ImportError:
            return None

        profile = self._profile_records(records, columns)
        prompt = {
            "task": task or "Choose outlier detection and imputation settings.",
            "known_domain": domain,
            "forced_method": method,
            "forced_imputation": imputation,
            "allowed_methods": sorted(self.tool.METHODS),
            "allowed_imputations": sorted(self.tool.IMPUTATIONS),
            "data_profile": profile,
            "context": self.context,
            "skills_md": self.skills[:6000],
        }

        try:
            client_kwargs = {"api_key": api_key}
            if endpoint:
                client_kwargs["base_url"] = endpoint
            client = OpenAI(**client_kwargs)
            response = client.chat.completions.create(
                model=model,
                temperature=temperature,
                response_format={"type": "json_object"},
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "You are the planning brain for an outlier agent. "
                            "Return only JSON with keys: method, imputation, "
                            "domain, reason. Choose only allowed values."
                        ),
                    },
                    {"role": "user", "content": json.dumps(prompt)},
                ],
            )
            content = response.choices[0].message.content or "{}"
            raw_plan = json.loads(content)
            if method:
                raw_plan["method"] = method
            if imputation:
                raw_plan["imputation"] = imputation
            raw_plan["source"] = "llm"
            return self._validated_plan(raw_plan)
        except Exception:
            return None

    def _validated_plan(self, raw_plan: dict[str, Any]) -> Plan:
        method = str(raw_plan.get("method", "")).lower().strip()
        imputation = str(raw_plan.get("imputation", "")).lower().strip()
        if method not in self.tool.METHODS:
            method = self.context["defaults"]["method"]
        if imputation not in self.tool.IMPUTATIONS:
            imputation = self.context["defaults"]["imputation"]

        domain = str(raw_plan.get("domain") or "general").lower().strip()
        reason = str(raw_plan.get("reason") or "Selected by Outlier_agent.")
        source = str(raw_plan.get("source") or "rules")
        return Plan(method, imputation, domain, reason, source)

    def _infer_domain(
        self, records: list[dict[str, Any]], columns: list[str] | None
    ) -> str:
        if not records:
            return "general"
        names = " ".join(columns or records[0].keys()).lower()
        if any(token in names for token in ("price", "revenue", "return", "asset")):
            return "finance"
        if any(token in names for token in ("patient", "blood", "heart", "dose")):
            return "healthcare"
        if any(token in names for token in ("sensor", "temperature", "humidity", "volt")):
            return "sensor"
        if any(token in names for token in ("defect", "tolerance", "batch", "yield")):
            return "manufacturing"
        if any(token in names for token in ("survey", "rating", "likert", "response")):
            return "survey"
        return "general"

    def _profile_records(
        self, records: list[dict[str, Any]], columns: list[str] | None
    ) -> dict[str, Any]:
        selected_columns = columns or list(records[0].keys()) if records else []
        numeric_columns = self.tool._resolve_numeric_columns(records, columns)
        preview = records[:5]
        return {
            "row_count": len(records),
            "selected_columns": selected_columns,
            "numeric_columns": numeric_columns,
            "preview_rows": preview,
        }


def parse_columns(raw_columns: str | None) -> list[str] | None:
    if not raw_columns:
        return None
    return [column.strip() for column in raw_columns.split(",") if column.strip()]


def load_env_file(path: str | Path = ENV_PATH) -> None:
    path = Path(path)
    if not path.exists():
        return

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("export "):
            line = line[len("export ") :].strip()
        if "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = strip_env_value(value.strip())
        if key and key not in os.environ:
            os.environ[key] = value


def strip_env_value(value: str) -> str:
    if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
        return value[1:-1]
    return value


def configured_secret(value: str | None) -> bool:
    if not value:
        return False
    normalized = value.strip().lower()
    return normalized not in {
        "replace-with-your-api-key",
        "your-api-key",
        "your_key_here",
        "changeme",
    }


def load_table(path: str | Path, *, sheet: str | None = None) -> list[dict[str, Any]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return load_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return load_xlsx(path, sheet=sheet)
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, list):
            return [dict(row) for row in payload]
        if isinstance(payload, dict) and isinstance(payload.get("records"), list):
            return [dict(row) for row in payload["records"]]
    raise ValueError("Supported input formats are .csv, .json, .xlsx, and .xlsm.")


def load_xlsx(path: Path, *, sheet: str | None = None) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Reading Excel files requires openpyxl. Install it or export the file as CSV."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.active
        rows = worksheet.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if not raw_headers:
            return []
        headers = unique_headers(raw_headers)
        records: list[dict[str, Any]] = []
        for row in rows:
            record: dict[str, Any] = {}
            has_value = False
            for header, value in zip_longest(headers, row, fillvalue=None):
                if header is None:
                    continue
                record[header] = "" if value is None else value
                has_value = has_value or value not in (None, "")
            if has_value:
                records.append(record)
        return records
    finally:
        workbook.close()


def unique_headers(raw_headers: tuple[Any, ...]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, raw_header in enumerate(raw_headers, start=1):
        base = str(raw_header).strip() if raw_header not in (None, "") else f"column_{index}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return headers


def demo_records() -> list[dict[str, Any]]:
    return [
        {"id": "1", "temperature": "21.1", "pressure": "101.2"},
        {"id": "2", "temperature": "21.5", "pressure": "101.0"},
        {"id": "3", "temperature": "21.3", "pressure": "100.9"},
        {"id": "4", "temperature": "97.2", "pressure": "101.1"},
        {"id": "5", "temperature": "21.4", "pressure": "100.8"},
        {"id": "6", "temperature": "21.2", "pressure": "160.0"},
    ]


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Detect and impute tabular outliers.")
    parser.add_argument(
        "input",
        nargs="?",
        help="Input .csv, .json, .xlsx, or .xlsm path. Omit to run demo data.",
    )
    parser.add_argument("--output", help="Cleaned CSV output path.")
    parser.add_argument("--report", help="JSON report output path.")
    parser.add_argument("--sheet", help="Excel sheet name. Defaults to the active sheet.")
    parser.add_argument("--columns", help="Comma-separated columns to inspect.")
    parser.add_argument("--domain", help="Optional domain profile, such as finance or sensor.")
    parser.add_argument(
        "--method",
        choices=sorted(OutlierTool.METHODS),
        help="Force a detection method.",
    )
    parser.add_argument(
        "--imputation",
        choices=sorted(OutlierTool.IMPUTATIONS),
        help="Force an imputation strategy.",
    )
    parser.add_argument("--task", help="Natural-language instruction for the LLM planner.")
    parser.add_argument("--no-llm", action="store_true", help="Disable LLM planning.")
    parser.add_argument(
        "--print-records",
        action="store_true",
        help="Print cleaned records to stdout as JSON.",
    )
    return parser


def main() -> None:
    args = build_parser().parse_args()
    records = load_table(args.input, sheet=args.sheet) if args.input else demo_records()
    columns = parse_columns(args.columns)
    agent = OutlierAgent(use_llm=not args.no_llm)
    result = agent.run(
        records,
        columns=columns,
        domain=args.domain,
        method=args.method,
        imputation=args.imputation,
        task=args.task,
    )

    input_path = Path(args.input) if args.input else None
    output_path = Path(args.output) if args.output else None
    report_path = Path(args.report) if args.report else None
    if input_path and output_path is None:
        output_path = input_path.with_name(f"{input_path.stem}_cleaned.csv")
    if report_path is None:
        if output_path:
            report_path = output_path.with_name(f"{output_path.stem}_report.json")
        else:
            report_path = Path("outlier_report.json")

    if output_path:
        save_csv(output_path, result["cleaned_records"])
    report_payload = {"plan": result["plan"], "report": result["report"]}
    report_path.write_text(json.dumps(report_payload, indent=2), encoding="utf-8")

    print(json.dumps(report_payload, indent=2))
    if args.print_records:
        print(json.dumps(result["cleaned_records"], indent=2))


if __name__ == "__main__":
    main()
